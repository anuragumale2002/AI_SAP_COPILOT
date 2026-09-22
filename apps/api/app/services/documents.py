import io
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass
class ExtractedChunk:
    page: int
    locator: str
    text: str


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _paragraph_chunks(text: str, page: int) -> list[ExtractedChunk]:
    paragraphs = [_clean(p) for p in re.split(r"\n\s*\n|(?<=\.)\s+(?=[A-Z])", text)]
    paragraphs = [p for p in paragraphs if len(p) >= 35]
    return [
        ExtractedChunk(page=page, locator=f"Page {page}, paragraph {i}", text=p[:1800])
        for i, p in enumerate(paragraphs, 1)
    ]


def _extract_docx(doc: DocxDocument) -> list[ExtractedChunk]:
    """Extract paragraphs *and tables* while retaining useful structural locators.

    BRDs frequently put requirements, mappings, test cases and screen fields in
    Word tables. Reading only ``doc.paragraphs`` silently loses that context.
    """
    chunks: list[ExtractedChunk] = []
    current_heading = ""
    paragraph_no = 0
    for paragraph in doc.paragraphs:
        text = _clean(paragraph.text)
        if not text:
            continue
        style_name = (getattr(paragraph.style, "name", "") or "").lower()
        if style_name.startswith("heading"):
            current_heading = text
        if len(text) < 20 and not style_name.startswith("heading"):
            continue
        paragraph_no += 1
        locator = f"Paragraph {paragraph_no}"
        if current_heading and current_heading != text:
            locator += f" under {current_heading[:80]}"
        chunks.append(ExtractedChunk(1, locator, text[:1800]))

    for table_no, table in enumerate(doc.tables, 1):
        for row_no, row in enumerate(table.rows, 1):
            cells = [_clean(cell.text) for cell in row.cells]
            if not any(cells):
                continue
            row_text = " | ".join(cells)
            if len(row_text) < 12:
                continue
            chunks.append(
                ExtractedChunk(
                    1,
                    f"Table {table_no}, row {row_no}",
                    row_text[:2400],
                )
            )
    return chunks


def extract_document(filename: str, content: bytes) -> tuple[int, list[ExtractedChunk]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        pages = PdfReader(io.BytesIO(content)).pages
        chunks = [
            chunk
            for page_no, page in enumerate(pages, 1)
            for chunk in _paragraph_chunks(page.extract_text() or "", page_no)
        ]
        return len(pages), chunks
    if suffix == ".docx":
        doc = DocxDocument(io.BytesIO(content))
        return 1, _extract_docx(doc)
    if suffix in {".txt", ".md"}:
        return 1, _paragraph_chunks(content.decode("utf-8", errors="replace"), 1)
    raise ValueError("Supported document types: PDF, DOCX, TXT, and Markdown")
