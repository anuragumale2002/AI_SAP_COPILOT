from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Spacer, Table, TableStyle
from reportlab.platypus.tableofcontents import TableOfContents

from .fsd_export import (
    DARK,
    GREEN,
    _FSDTemplate,
    _draw_logos,
    _draw_process_flow,
    _draw_process_legend,
    _overlay_pdf_chrome,
    _pdf_bullets,
    _pdf_picture,
    _pdf_styles,
    _pdf_table,
    _pdf_table_number,
    _pdf_text,
    _safe,
    _shade,
    _short_text,
    _style_word,
    _word_caption,
    _word_list,
    _word_picture,
    _word_table,
    _word_toc,
)
from .project_identity import clean_document_title, looks_like_person_name, process_identifier_from_name


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "project"


def _center(doc: WordDocument, text: str, size: float, color: str, bold: bool = True, space_after: int = 6) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def _tdd_display_name(data: dict, fallback: str) -> str:
    info = data.get("document_information") if isinstance(data.get("document_information"), dict) else {}
    return clean_document_title(
        info.get("project") or data.get("title"),
        fallback or "SAP Business Process",
    )


def _tdd_process_identifier(data: dict, display_name: str) -> str:
    info = data.get("document_information") if isinstance(data.get("document_information"), dict) else {}
    identifier = str(info.get("process_identifier") or "")
    if identifier and not looks_like_person_name(identifier):
        return identifier
    project_value = str(info.get("project") or "")
    if project_value and not looks_like_person_name(project_value) and re.search(r"[A-Z0-9]+-[A-Z0-9]+", project_value):
        return project_value
    return process_identifier_from_name(display_name)


def _apply_tdd_identity(data: dict, fallback: str) -> str:
    display_name = _tdd_display_name(data, fallback)
    data["title"] = display_name
    info = data.setdefault("document_information", {})
    info["project"] = display_name
    info["process_identifier"] = _tdd_process_identifier(data, display_name)
    source = data.get("source_fsd") if isinstance(data.get("source_fsd"), dict) else {}
    if source:
        source["title"] = clean_document_title(source.get("title"), display_name)
        data["source_fsd"] = source
    return display_name


def _word_screenshot_placeholder(doc: WordDocument, label: str, instruction: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _shade(cell, "FFF4E5")
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "dashed")
        node.set(qn("w:sz"), "12")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "C47A2C")
        borders.append(node)
    properties.append(borders)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), "1600")
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = paragraph.add_run("SCREENSHOT REQUIRED")
    heading.bold = True
    heading.font.size = Pt(11)
    heading.font.color.rgb = RGBColor.from_string("8A5A12")
    detail = cell.add_paragraph(instruction)
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in detail.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("5B6762")
    _word_caption(doc, label)


def _pdf_screenshot_placeholder(styles: dict, label: str, instruction: str) -> list:
    box = Table(
        [[_pdf_text(f"SCREENSHOT REQUIRED\n{instruction}", styles["Body"])]],
        colWidths=[6.5 * inch],
    )
    box.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fff4e5")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#c47a2c")),
        ("TOPPADDING", (0, 0), (-1, -1), 18),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 18),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    return [Spacer(1, 6), box, _pdf_text(label, styles["Caption"])]


def _pdf_placeholder_blocks(styles: dict, items: list[tuple[str, str]]) -> list:
    parts: list = []
    for label, instruction in items:
        parts.extend(_pdf_screenshot_placeholder(styles, label, instruction))
    return parts


def _screen_placeholders(screens: list) -> list[tuple[str, str]]:
    if not screens:
        return [(
            "Figure: Primary Fiori application screenshot",
            "Insert a screenshot of the primary Fiori application (list report or object page) showing search, key fields and actions as they appear to the business user.",
        )]
    placeholders = []
    for screen in screens[:6]:
        name = _safe(screen.get("name") or "Fiori screen")
        technology = _safe(screen.get("technology") or "SAP Fiori")
        placeholders.append((
            f"Figure: {name} screenshot",
            f"Insert a screenshot of '{name}' ({technology}). Capture the layout, key fields/controls and actions a reviewer would see from the Fiori launchpad.",
        ))
    return placeholders


def _cover(doc: WordDocument, data: dict, project_name: str) -> None:
    info = data.get("document_information", {})
    source = data.get("source_fsd") if isinstance(data.get("source_fsd"), dict) else {}
    _center(doc, "TECHNICAL DESIGN SPECIFICATION", 22, DARK, space_after=14)
    _center(doc, _safe(data.get("title") or project_name).upper(), 16, GREEN, space_after=8)
    _center(doc, _safe(info.get("application")), 13, DARK, bold=False, space_after=10)
    _center(doc, _safe(data.get("generated_on")), 12, DARK, bold=False, space_after=18)
    _word_table(doc, ["Document Attribute", "Value"], [
        ["Application", info.get("application")],
        ["Project / Process Identifier", info.get("process_identifier") or project_name],
        ["Process Area", info.get("process_area")],
        ["Version", info.get("version") or "0.1"],
        ["Status", data.get("status")],
        ["Prepared By", info.get("prepared_by") or "SAP Project Copilot"],
        ["Source FSD", f"{_safe(source.get('title'))}, Version {_safe(source.get('version'))}"],
    ], [2.2, 4.4], title="Cover document attributes")
    _center(doc, f"Prepared by: {_safe(info.get('prepared_by') or 'SAP Project Copilot')} (Draft)", 11, DARK, bold=False, space_after=4)
    _center(
        doc,
        f"Source FSD: {_safe(source.get('title'))}, Version {_safe(source.get('version'))}",
        11,
        DARK,
        bold=False,
        space_after=16,
    )

    doc.add_heading("Revision History", 1)
    _word_table(
        doc,
        ["Version No.", "Date", "Author", "Revision Description"],
        [[item.get("version"), item.get("date"), item.get("author"), item.get("description")] for item in data.get("revision_history", [])],
        [.85, 1.1, 1.55, 3.1], title="Document history",
    )
    doc.add_heading("Authorisers for sign off", 2)
    _word_table(
        doc,
        ["Role", "Name", "Date Signed Off"],
        [[item.get("role"), item.get("name"), item.get("date") or item.get("status") or "TBD"] for item in data.get("sign_offs", [])],
        [2.3, 2.2, 2.1], title="Authorisers for sign off",
    )
    doc.add_heading("Distribution - in addition to Authorisers", 2)
    _word_table(
        doc,
        ["Name / Team", "Role"],
        [[item.get("name"), item.get("role")] for item in data.get("distribution", [])],
        [3.3, 3.3], title="Document distribution",
    )
    contents = doc.add_paragraph("Contents")
    for run in contents.runs:
        run.font.name, run.font.size, run.font.bold = "Calibri", Pt(16), True
        run.font.color.rgb = RGBColor.from_string(GREEN)
    contents.paragraph_format.space_after = Pt(10)
    _word_toc(doc)
    doc.add_page_break()


def _add_contents(doc: WordDocument, data: dict, images: dict[str, Path]) -> None:
    info = data.get("document_information", {})
    ops = data.get("operational_considerations") if isinstance(data.get("operational_considerations"), dict) else {}

    doc.add_heading("1. Document Information", 1)
    doc.add_heading("1.1 General Data", 2)
    _word_table(doc, ["Information", "Detail"], [
        ["Extension ID", info.get("extension_id")],
        ["Extension Description", info.get("extension_description")],
        ["Application", info.get("application")],
        ["Process Area", info.get("process_area")],
        ["Project", info.get("project")],
        ["Process Identifier", info.get("process_identifier")],
        ["Job Name", info.get("job_name")],
        ["Functional Designer", info.get("functional_designer")],
        ["Technical Designer", info.get("technical_designer")],
        ["Developer", info.get("developer")],
        ["Tools / Technology", info.get("tools_technology")],
    ], [2.15, 4.45], title="General data")
    doc.add_heading("1.2 Related Documents", 2)
    _word_table(
        doc,
        ["Document", "Location / Reference"],
        [[item.get("document"), item.get("location")] for item in data.get("related_documents", [])],
        [2.6, 4.0], title="Related documents",
    )

    doc.add_heading("2. Development Overview", 1)
    doc.add_heading("2.1 Requirements Summary", 2)
    doc.add_paragraph(_safe(data.get("requirements_summary")))
    _word_table(
        doc,
        ["Capability", "Technical Response", "Key Requirement IDs"],
        [[item.get("capability"), item.get("technical_response"), item.get("requirement_ids")] for item in data.get("capability_map", [])],
        [1.7, 3.15, 1.75], title="Requirements summary",
    )
    doc.add_heading("2.2 Assumptions", 2)
    _word_list(doc, data.get("assumptions", []))
    doc.add_heading("2.3 Dependencies/Constraints", 2)
    _word_list(doc, data.get("dependencies", []))
    doc.add_heading("2.4 Objects or Transaction Affected", 2)
    _word_table(
        doc,
        ["Object / Application", "Impact"],
        [[f"{item.get('object_type')}: {item.get('object_name')}", item.get("impact") or item.get("purpose")] for item in data.get("affected_objects", [])],
        [2.4, 4.2], title="Objects or transactions affected",
    )
    doc.add_heading("2.5 Standards", 2)
    _word_list(doc, data.get("standards", []))
    doc.add_heading("2.6 Security, Integrity and Controls", 2)
    _word_table(
        doc,
        ["Control", "Technical Design"],
        [[item.get("control"), item.get("design")] if isinstance(item, dict) else [item, "Confirm during design review"] for item in data.get("security_controls", [])],
        [2.0, 4.6], title="Security, integrity and controls",
    )
    doc.add_heading("2.7 Error Handling / Messages", 2)
    _word_table(
        doc,
        ["ID", "Scenario", "User Message / Behavior", "Technical Handling"],
        [[item.get("id"), item.get("scenario"), item.get("user_message"), item.get("technical_handling")] for item in data.get("error_handling", [])],
        [.7, 1.45, 2.25, 2.2], title="Error handling messages",
    )

    doc.add_heading("3. Detailed Technical Specifications", 1)
    doc.add_heading("3.1 Technical Flow Diagram", 2)
    if images.get("process"):
        _word_picture(doc, images["process"])
        _word_caption(doc, "Figure 1: Technical flow derived from the approved Functional Specification")
    if images.get("legend"):
        _word_picture(doc, images["legend"], max_height=4.2)
        _word_caption(doc, "Figure 2: Process diagram legend")
    _word_screenshot_placeholder(
        doc,
        "Figure: As-built technical sequence screenshot",
        "Insert a screenshot of the as-built Fiori-to-S/4HANA technical sequence (SAP Signavio, Solution Manager, or architect diagram) if it differs from the generated flow above.",
    )
    doc.add_paragraph("Description of flow diagram:", style="Record Heading")
    _word_list(doc, [
        f"{_safe(step.get('step_id'))}: {_short_text(step.get('system_behavior') or step.get('activity'), 220)} Outcome: {_short_text(step.get('outcome'), 120)}"
        for step in data.get("process_steps", [])
    ])
    doc.add_heading("3.2 Fiori Technical Flow Description", 2)
    _word_table(
        doc,
        ["Flow", "UI / Service Behavior", "Method / Mechanism"],
        [[item.get("flow"), item.get("behavior"), item.get("mechanism")] for item in data.get("fiori_flows", [])],
        [1.55, 3.25, 1.8], title="Fiori technical flow",
    )
    doc.add_heading("3.3 Selection Screen Details", 2)
    _word_table(
        doc,
        ["Screen / View", "Proposed Technology", "Key Fields / Controls", "Actions"],
        [[item.get("name"), item.get("technology"), item.get("fields"), item.get("actions")] for item in data.get("screens", [])]
        or [["TBD", "SAP Fiori", "Screen names, controls and value helps are finalized during UI design.", "TBD"]],
        [1.45, 1.55, 2.15, 1.45], title="Selection screen details",
    )
    for label, instruction in _screen_placeholders(data.get("screens") or []):
        _word_screenshot_placeholder(doc, label, instruction)
    doc.add_heading("3.4 Security and Authorization", 2)
    roles = data.get("roles_and_authorizations") or []
    _word_table(
        doc,
        ["Role", "Data Scope", "Technical Authorization Behavior", "SoD / Control"],
        [[
            item.get("business_role") or item.get("role_id") or item.get("role"),
            item.get("data_scope") or "Authorized business records only",
            item.get("authorization_object") or item.get("activities") or item.get("transaction_or_app") or "Backend authorization on every protected operation",
            "Least privilege; UI visibility is not an authorization control",
        ] for item in roles] or [["Business User", "Authorized records", "Backend checks for every protected operation", "Least privilege"]],
        [1.35, 1.55, 2.2, 1.5], title="Security and authorization",
    )
    _word_screenshot_placeholder(
        doc,
        "Figure: Business role / authorization screenshot",
        "Insert a screenshot of the PFCG business role (or equivalent IAM assignment) showing authorization objects and organizational restrictions for the primary process role.",
    )
    doc.add_heading("3.5 Processing and Operational Considerations", 2)
    doc.add_heading("3.5.1 Dependencies", 3)
    _word_list(doc, ops.get("dependencies") or data.get("dependencies", []))
    doc.add_heading("3.5.2 Re-Use Details", 3)
    _word_list(doc, ops.get("reuse") or ["Reuse standard SAP Fiori, OData, authentication and authorization capabilities where they satisfy the FSD"])
    doc.add_heading("3.5.3 Fiori Application HTML5/JSP/CSS Details (Customized Objects Only)", 3)
    _word_table(doc, ["Item", "Design"], ops.get("fiori_html") or [["Application type", "Prefer Fiori elements List Report + Object Page"]], [2.1, 4.5], title="Fiori HTML5 / CSS design")
    doc.add_heading("3.5.4 Business Server Pages / Extensions", 3)
    doc.add_paragraph(_safe(ops.get("bsp") or "N/A for the baseline design. No custom BSP application is required by the FSD."))
    doc.add_heading("3.5.5 Git Repository Details", 3)
    _word_table(doc, ["Item", "Detail"], ops.get("git") or [["Git repository", "TBD"]], [2.4, 4.2], title="Git repository details")
    doc.add_heading("3.5.6 Multi-Site Details", 3)
    doc.add_paragraph(_safe(ops.get("multi_site") or "Confirm whether one global template or site-specific configuration is required."))
    doc.add_heading("3.5.7 Other", 3)
    _word_list(doc, ops.get("other") or ["Protect state-changing calls with standard SAP OData CSRF/session mechanisms"])

    doc.add_heading("4. Technical Requirements", 1)
    doc.add_heading("4.1 Database Tables", 2)
    mappings = data.get("data_mappings") or []
    _word_table(
        doc,
        ["Business Data", "Source / Contract", "Technical Rule"],
        [[
            item.get("business_field") or item.get("mapping_id") or item.get("name") or "Business field",
            item.get("source") or item.get("sap_field") or item.get("source_object") or "Approved SAP persistence / OData",
            item.get("rule") or item.get("transformation") or item.get("description") or "Backend remains authoritative",
        ] for item in mappings] or [["TBD", "Approved SAP persistence / OData", "Final tables, CDS views and entity sets remain TBD until landscape confirmation"]],
        [1.9, 2.35, 2.35], title="Database and data mapping",
    )
    doc.add_heading("4.2 External Programs", 2)
    interfaces = data.get("interfaces") or []
    _word_table(
        doc,
        ["Program / Interface", "Purpose", "Direction"],
        [[
            item.get("name") or item.get("interface_id") or "Interface",
            item.get("purpose") or item.get("description") or "Confirm during landscape review",
            item.get("direction") or item.get("pattern") or "TBD",
        ] for item in interfaces] or [["None identified in the FSD", "No external program is invented in this draft", "N/A"]],
        [2.0, 3.1, 1.5], title="External programs",
    )
    doc.add_heading("4.3 Development Information", 2)
    _word_table(doc, ["Attribute", "Value"], [
        ["Program ID", "TBD"],
        ["Program Type", info.get("program_type")],
        ["Module / Process Area", info.get("process_area")],
        ["Development Class", "TBD"],
        ["Message Class", "TBD"],
        ["Tools / Technology", info.get("tools_technology")],
    ], [2.15, 4.45], title="Development information")
    doc.add_heading("4.4 Transport Number(s)", 2)
    _word_table(doc, ["Item", "Detail"], [["Workbench request", "TBD"], ["Customizing request", "TBD"], ["Package / development class", "TBD"]], [2.4, 4.2], title="Transport requests")
    doc.add_heading("4.5 Detailed Design", 2)
    doc.add_heading("4.5.1 SAP Objects", 3)
    _word_table(
        doc,
        ["Object Type", "Object Name", "Purpose"],
        [[item.get("object_type"), item.get("object_name"), item.get("purpose")] for item in data.get("affected_objects", [])],
        [1.7, 1.9, 3.0], title="SAP objects",
    )
    _word_screenshot_placeholder(
        doc,
        "Figure: SAP repository objects screenshot",
        "Insert a screenshot of the assigned repository objects (ADT/SE80 package, Fiori app, OData service, ABAP class) once technical names are confirmed.",
    )
    doc.add_heading("4.5.2 Request Mechanism / Trigger", 3)
    doc.add_paragraph("SAP Fiori action through an approved OData service unless fit-to-standard analysis selects a standard application or API. Workflow decisions, where required by the FSD, execute through SAP workflow / My Inbox services rather than direct UI status mutation.")
    doc.add_heading("4.5.3 Pseudocode", 3)
    for item in data.get("components", []):
        doc.add_paragraph(
            f"{_safe(item.get('requirement_key'))} - Validate authorization and input; execute {_safe(item.get('title'))}; "
            f"commit only after all validations pass; return {_short_text(item.get('validation'), 160)}.",
            style="List Number",
        )
    doc.add_heading("4.5.3.1 Custom Table Development", 3)
    doc.add_paragraph("No custom table is invented in this draft. Persistence remains in approved SAP objects unless a documented fit-gap requires otherwise.")
    doc.add_heading("4.5.3.2 User Exit", 3)
    doc.add_paragraph("No user exit is proposed until a confirmed standard-process gap is recorded.")
    doc.add_heading("4.5.3.3 Field Exit", 3)
    doc.add_paragraph("N/A. Field exits are not part of the baseline Fiori / OData design.")
    doc.add_heading("4.5.3.4 Technical Design Constraints", 3)
    _word_list(doc, [
        "Do not invent unapproved object, service, package or transport names",
        "Do not hard-code system URLs, credentials or landscape identifiers",
        "UI visibility is not a substitute for backend authorization",
        "Object names marked TBD require technical lead assignment before build",
    ])
    doc.add_heading("4.5.3.5 Programmers Notes", 3)
    _word_list(doc, [
        "Select only fields required for the current screen and business role",
        "Map backend errors to controlled business messages; retain technical detail in support logs only",
        "Include a request/correlation identifier and business key in support logging where available",
        "Use semantic status controls so meaning is accessible without color",
    ])
    doc.add_heading("4.5.4 OData and Data Mapping Design", 3)
    _word_table(
        doc,
        ["Business Data", "Source / Contract", "UI Target", "Technical Rule"],
        [[
            item.get("business_field") or item.get("mapping_id") or "Business data",
            item.get("source") or item.get("sap_field") or "Approved OData / CDS",
            item.get("ui_target") or item.get("screen") or "Fiori worklist / object page",
            item.get("rule") or item.get("transformation") or "GET for retrieval; POST/PATCH for permitted change; backend authorization",
        ] for item in mappings] or [["Approved FSD data", "OData service TBD", "Fiori application", "Backend remains authoritative"]],
        [1.5, 1.7, 1.55, 1.85], title="OData and data mapping",
    )
    _word_screenshot_placeholder(
        doc,
        "Figure: OData service metadata screenshot",
        "Insert a screenshot of the OData service metadata (SEGW, /IWFND/MAINT_SERVICE, or $metadata) for the approved entity set and operations.",
    )
    doc.add_heading("4.5.5 Workflow Technical Design", 3)
    variants = data.get("process_steps") or []
    _word_table(
        doc,
        ["Rule / Condition", "Technical Behavior", "Open Decision"],
        [[
            item.get("activity") or item.get("step_id"),
            item.get("system_behavior") or item.get("outcome"),
            "Confirm agent determination, thresholds and exception handling during design review" if "approv" in _safe(item.get("activity")).lower() or "approv" in _safe(item.get("system_behavior")).lower() else "None once configuration is approved",
        ] for item in variants[:10]] or [["Workflow (if required by FSD)", "Use standard SAP workflow / My Inbox when the FSD specifies approval", "Confirm during design review"]],
        [1.8, 2.7, 2.1], title="Workflow technical design",
    )
    _word_screenshot_placeholder(
        doc,
        "Figure: Flexible Workflow / My Inbox screenshot",
        "Insert a screenshot of the Flexible Workflow scenario or My Inbox task configuration once agent determination and thresholds are confirmed.",
    )
    doc.add_heading("4.5.6 Status Model", 3)
    statuses = data.get("status_definitions") or []
    _word_table(
        doc,
        ["Status", "Source", "UI Semantic State", "Definition"],
        [[
            item.get("status") or item.get("name") or item.get("status_id"),
            item.get("source") or item.get("domain") or "SAP backend",
            item.get("semantic_state") or item.get("ui_state") or "TBD",
            item.get("definition") or item.get("description") or "Confirm during design review",
        ] for item in statuses] or [["Draft / In Process / Complete", "SAP backend", "None / Information / Success", "Business-facing states derived from the FSD; the UI must not mutate status directly"]],
        [1.45, 1.35, 1.45, 2.35], title="Status model",
    )
    doc.add_heading("4.5.7 Performance and Monitoring Design", 3)
    nfr = [_safe(item.get("requirement") or item.get("description") or item) for item in data.get("nfr") or []]
    _word_list(doc, nfr or [
        "Use server-side filter, sort, page and select. Avoid loading the full population into browser memory.",
        "Instrument service latency, failures and authorization denials for support without exposing secrets.",
        "Measure performance with representative dataset sizes and organizational authorization filters in the target landscape.",
    ])

    doc.add_heading("5. Testing Requirements", 1)
    doc.add_heading("5.1 Key Unit / Assembly Test Conditions", 2)
    _word_table(
        doc,
        ["ID", "Condition / Test", "Expected Result", "Req Ref."],
        [[item.get("id"), item.get("condition"), item.get("expected_result"), item.get("requirement_key") or item.get("cycle_ref")] for item in data.get("test_conditions", [])],
        [.7, 2.25, 2.2, 1.45], title="Unit and assembly test conditions",
    )
    _word_screenshot_placeholder(
        doc,
        "Figure: Unit / assembly test evidence screenshot",
        "Insert screenshots of unit or assembly test evidence (ABAP Unit, Gateway client, or Fiori test run) for the critical path covered above.",
    )

    doc.add_heading("6. Outstanding Issues", 1)
    _word_table(
        doc,
        ["Issue No.", "Description", "Assigned To", "Status", "Impact", "Resolution / Decision Needed"],
        [[item.get("issue_no"), item.get("description"), item.get("assigned_to"), item.get("status"), item.get("impact"), item.get("resolution")] for item in data.get("outstanding_issues", [])],
        [.75, 1.7, 1.05, .7, .85, 1.55], title="Outstanding issues",
    )

    doc.add_heading("7. Appendix", 1)
    doc.add_heading("7.1 Glossary of Terms", 2)
    _word_table(
        doc,
        ["Term", "Definition"],
        [[item.get("term"), item.get("definition")] for item in data.get("glossary", [])],
        [1.7, 4.9], title="Glossary of terms",
    )
    doc.add_heading("7.2 Additional Supporting / Reference Documentation", 2)
    _word_list(doc, [
        f"Functional Solution Design: {_safe((data.get('source_fsd') or {}).get('title'))}, Version {_safe((data.get('source_fsd') or {}).get('version'))}",
        "Approved BRD and generated requirement baseline",
        "Quality Test Pack, traceability matrix, and generated Fiori/ABAP starter-code package",
        "Target system OData service metadata, API inventory, transports and package names - TBD",
    ])
    doc.add_heading("7.3 Key Requirement Traceability", 2)
    _word_table(
        doc,
        ["Design Area", "FSD Requirements Covered", "Technical Design Sections"],
        [[item.get("capability"), item.get("requirement_ids"), "2.1, 3.2, 3.3, 4.5"] for item in data.get("traceability") or data.get("capability_map", [])],
        [1.9, 2.5, 2.2], title="Requirement traceability",
    )
    doc.add_heading("7.4 Technical Design Review Notes", 2)
    _word_list(doc, data.get("review_notes", []))


def _build_tdd_pdf(data: dict, images: dict[str, Path], output: Path, project_name: str) -> None:
    _pdf_table_number.set(0)
    styles = _pdf_styles()
    info = data.get("document_information", {})
    source = data.get("source_fsd") if isinstance(data.get("source_fsd"), dict) else {}
    ops = data.get("operational_considerations") if isinstance(data.get("operational_considerations"), dict) else {}
    toc = TableOfContents()
    toc.levelStyles = [styles["TOC0"], styles["TOC1"]]
    toc.dotsMinLevel = 0
    story: list = [
        Spacer(1, .35 * inch),
        _pdf_text("TECHNICAL DESIGN SPECIFICATION", styles["CoverTitle"]),
        _pdf_text(data.get("title") or project_name, styles["CoverSub"]),
        _pdf_text(info.get("application"), styles["CoverSub"]),
        _pdf_text(data.get("generated_on"), styles["Body"]),
        _pdf_table(["Document Attribute", "Value"], [
            ["Application", info.get("application")],
            ["Project / Process Identifier", info.get("process_identifier") or project_name],
            ["Process Area", info.get("process_area")],
            ["Version", info.get("version") or "0.1"],
            ["Status", data.get("status")],
            ["Prepared By", info.get("prepared_by") or "SAP Project Copilot"],
            ["Source FSD", f"{_safe(source.get('title'))}, Version {_safe(source.get('version'))}"],
        ], [2.2, 4.3], styles, title="Cover document attributes"),
        _pdf_text("Revision History", styles["FSDH1"]),
        _pdf_table(["Version No.", "Date", "Author", "Revision Description"],
                   [[item.get("version"), item.get("date"), item.get("author"), item.get("description")] for item in data.get("revision_history", [])],
                   [.85, 1.05, 1.5, 3.1], styles, title="Document history"),
        _pdf_text("Authorisers for sign off", styles["FSDH2"]),
        _pdf_table(["Role", "Name", "Date Signed Off"],
                   [[item.get("role"), item.get("name"), item.get("date") or item.get("status") or "TBD"] for item in data.get("sign_offs", [])],
                   [2.2, 2.15, 2.15], styles, title="Authorisers for sign off"),
        _pdf_text("Distribution - in addition to Authorisers", styles["FSDH2"]),
        _pdf_table(["Name / Team", "Role"], [[item.get("name"), item.get("role")] for item in data.get("distribution", [])],
                   [3.25, 3.25], styles, title="Document distribution"),
        PageBreak(), _pdf_text("Contents", styles["CoverSub"]), Spacer(1, 8), toc, PageBreak(),
        _pdf_text("1. Document Information", styles["FSDH1"]),
        _pdf_text("1.1 General Data", styles["FSDH2"]),
        _pdf_table(["Information", "Detail"], [
            ["Extension ID", info.get("extension_id")], ["Extension Description", info.get("extension_description")],
            ["Application", info.get("application")], ["Process Area", info.get("process_area")],
            ["Project", info.get("project")], ["Process Identifier", info.get("process_identifier")],
            ["Job Name", info.get("job_name")],
            ["Functional Designer", info.get("functional_designer")], ["Technical Designer", info.get("technical_designer")],
            ["Developer", info.get("developer")], ["Tools / Technology", info.get("tools_technology")],
        ], [2.1, 4.4], styles, title="General data"),
        _pdf_text("1.2 Related Documents", styles["FSDH2"]),
        _pdf_table(["Document", "Location / Reference"], [[item.get("document"), item.get("location")] for item in data.get("related_documents", [])],
                   [2.55, 3.95], styles, title="Related documents"),
        _pdf_text("2. Development Overview", styles["FSDH1"]),
        _pdf_text("2.1 Requirements Summary", styles["FSDH2"]),
        _pdf_text(data.get("requirements_summary"), styles["Body"]),
        _pdf_table(["Capability", "Technical Response", "Key Requirement IDs"],
                   [[item.get("capability"), item.get("technical_response"), item.get("requirement_ids")] for item in data.get("capability_map", [])],
                   [1.65, 3.15, 1.7], styles, title="Requirements summary"),
        _pdf_text("2.2 Assumptions", styles["FSDH2"]),
        *_pdf_bullets(data.get("assumptions", []), styles),
        _pdf_text("2.3 Dependencies/Constraints", styles["FSDH2"]),
        *_pdf_bullets(data.get("dependencies", []), styles),
        _pdf_text("2.4 Objects or Transaction Affected", styles["FSDH2"]),
        _pdf_table(["Object / Application", "Impact"],
                   [[f"{item.get('object_type')}: {item.get('object_name')}", item.get("impact") or item.get("purpose")] for item in data.get("affected_objects", [])],
                   [2.35, 4.15], styles, title="Objects or transactions affected"),
        _pdf_text("2.5 Standards", styles["FSDH2"]),
        *_pdf_bullets(data.get("standards", []), styles),
        _pdf_text("2.6 Security, Integrity and Controls", styles["FSDH2"]),
        _pdf_table(["Control", "Technical Design"],
                   [[item.get("control"), item.get("design")] if isinstance(item, dict) else [item, "Confirm during design review"] for item in data.get("security_controls", [])],
                   [2.0, 4.5], styles, title="Security, integrity and controls"),
        _pdf_text("2.7 Error Handling / Messages", styles["FSDH2"]),
        _pdf_table(["ID", "Scenario", "User Message / Behavior", "Technical Handling"],
                   [[item.get("id"), item.get("scenario"), item.get("user_message"), item.get("technical_handling")] for item in data.get("error_handling", [])],
                   [.65, 1.4, 2.25, 2.2], styles, title="Error handling messages"),
        _pdf_text("3. Detailed Technical Specifications", styles["FSDH1"]),
        _pdf_text("3.1 Technical Flow Diagram", styles["FSDH2"]),
    ]
    if images.get("process"):
        story += [_pdf_picture(images["process"]), _pdf_text("Figure 1: Technical flow derived from the approved Functional Specification", styles["Caption"])]
    if images.get("legend"):
        story += [_pdf_picture(images["legend"], max_height=4.0 * inch), _pdf_text("Figure 2: Process diagram legend", styles["Caption"])]
    story += _pdf_screenshot_placeholder(
        styles,
        "Figure: As-built technical sequence screenshot",
        "Insert a screenshot of the as-built Fiori-to-S/4HANA technical sequence (SAP Signavio, Solution Manager, or architect diagram) if it differs from the generated flow above.",
    )
    story += [
        _pdf_text("3.2 Fiori Technical Flow Description", styles["FSDH2"]),
        _pdf_table(["Flow", "UI / Service Behavior", "Method / Mechanism"],
                   [[item.get("flow"), item.get("behavior"), item.get("mechanism")] for item in data.get("fiori_flows", [])],
                   [1.5, 3.2, 1.8], styles, title="Fiori technical flow"),
        _pdf_text("3.3 Selection Screen Details", styles["FSDH2"]),
        _pdf_table(["Screen / View", "Proposed Technology", "Key Fields / Controls", "Actions"],
                   [[item.get("name"), item.get("technology"), item.get("fields"), item.get("actions")] for item in data.get("screens", [])]
                   or [["TBD", "SAP Fiori", "Finalized during UI design", "TBD"]],
                   [1.4, 1.5, 2.15, 1.45], styles, title="Selection screen details"),
    ]
    story += _pdf_placeholder_blocks(styles, _screen_placeholders(data.get("screens") or []))
    story += [
        _pdf_text("3.4 Security and Authorization", styles["FSDH2"]),
        _pdf_table(["Role", "Data Scope", "Technical Authorization Behavior", "SoD / Control"],
                   [[item.get("business_role") or item.get("role_id") or item.get("role"),
                     item.get("data_scope") or "Authorized business records only",
                     item.get("authorization_object") or item.get("activities") or "Backend authorization on every protected operation",
                     "Least privilege; UI visibility is not an authorization control"]
                    for item in (data.get("roles_and_authorizations") or [])]
                   or [["Business User", "Authorized records", "Backend checks for every protected operation", "Least privilege"]],
                   [1.3, 1.5, 2.2, 1.5], styles, title="Security and authorization"),
    ]
    story += _pdf_screenshot_placeholder(
        styles,
        "Figure: Business role / authorization screenshot",
        "Insert a screenshot of the PFCG business role (or equivalent IAM assignment) showing authorization objects and organizational restrictions for the primary process role.",
    )
    story += [
        _pdf_text("3.5 Processing and Operational Considerations", styles["FSDH2"]),
        _pdf_text("3.5.1 Dependencies", styles["FSDH3"]),
        *_pdf_bullets(ops.get("dependencies") or data.get("dependencies", []), styles),
        _pdf_text("3.5.2 Re-Use Details", styles["FSDH3"]),
        *_pdf_bullets(ops.get("reuse") or ["Reuse standard SAP Fiori, OData, authentication and authorization capabilities"], styles),
        _pdf_text("3.5.3 Fiori Application HTML5/JSP/CSS Details (Customized Objects Only)", styles["FSDH3"]),
        _pdf_table(["Item", "Design"], ops.get("fiori_html") or [["Application type", "Prefer Fiori elements List Report + Object Page"]],
                   [2.05, 4.45], styles, title="Fiori HTML5 / CSS design"),
        _pdf_text("3.5.4 Business Server Pages / Extensions", styles["FSDH3"]),
        _pdf_text(ops.get("bsp") or "N/A for the baseline design.", styles["Body"]),
        _pdf_text("3.5.5 Git Repository Details", styles["FSDH3"]),
        _pdf_table(["Item", "Detail"], ops.get("git") or [["Git repository", "TBD"]],
                   [2.35, 4.15], styles, title="Git repository details"),
        _pdf_text("3.5.6 Multi-Site Details", styles["FSDH3"]),
        _pdf_text(ops.get("multi_site") or "Confirm whether one global template or site-specific configuration is required.", styles["Body"]),
        _pdf_text("3.5.7 Other", styles["FSDH3"]),
        *_pdf_bullets(ops.get("other") or ["Protect state-changing calls with standard SAP OData CSRF/session mechanisms"], styles),
        _pdf_text("4. Technical Requirements", styles["FSDH1"]),
        _pdf_text("4.1 Database Tables", styles["FSDH2"]),
        _pdf_table(["Business Data", "Source / Contract", "Technical Rule"],
                   [[item.get("business_field") or item.get("mapping_id") or "Business field",
                     item.get("source") or item.get("sap_field") or "Approved SAP persistence / OData",
                     item.get("rule") or item.get("transformation") or "Backend remains authoritative"]
                    for item in (data.get("data_mappings") or [])]
                   or [["TBD", "Approved SAP persistence / OData", "Final tables remain TBD until landscape confirmation"]],
                   [1.85, 2.3, 2.35], styles, title="Database and data mapping"),
        _pdf_text("4.2 External Programs", styles["FSDH2"]),
        _pdf_table(["Program / Interface", "Purpose", "Direction"],
                   [[item.get("name") or item.get("interface_id") or "Interface",
                     item.get("purpose") or item.get("description") or "Confirm during landscape review",
                     item.get("direction") or item.get("pattern") or "TBD"]
                    for item in (data.get("interfaces") or [])]
                   or [["None identified in the FSD", "No external program is invented in this draft", "N/A"]],
                   [1.95, 3.05, 1.5], styles, title="External programs"),
        _pdf_text("4.3 Development Information", styles["FSDH2"]),
        _pdf_table(["Attribute", "Value"], [
            ["Program ID", "TBD"], ["Program Type", info.get("program_type")], ["Module / Process Area", info.get("process_area")],
            ["Development Class", "TBD"], ["Message Class", "TBD"], ["Tools / Technology", info.get("tools_technology")],
        ], [2.1, 4.4], styles, title="Development information"),
        _pdf_text("4.4 Transport Number(s)", styles["FSDH2"]),
        _pdf_table(["Item", "Detail"], [["Workbench request", "TBD"], ["Customizing request", "TBD"], ["Package / development class", "TBD"]],
                   [2.35, 4.15], styles, title="Transport requests"),
        _pdf_text("4.5 Detailed Design", styles["FSDH2"]),
        _pdf_text("4.5.1 SAP Objects", styles["FSDH3"]),
        _pdf_table(["Object Type", "Object Name", "Purpose"],
                   [[item.get("object_type"), item.get("object_name"), item.get("purpose")] for item in data.get("affected_objects", [])],
                   [1.65, 1.85, 3.0], styles, title="SAP objects"),
    ]
    story += _pdf_screenshot_placeholder(
        styles,
        "Figure: SAP repository objects screenshot",
        "Insert a screenshot of the assigned repository objects (ADT/SE80 package, Fiori app, OData service, ABAP class) once technical names are confirmed.",
    )
    story += [
        _pdf_text("4.5.2 Request Mechanism / Trigger", styles["FSDH3"]),
        _pdf_text("SAP Fiori action through an approved OData service unless fit-to-standard analysis selects a standard application or API.", styles["Body"]),
        _pdf_text("4.5.4 OData and Data Mapping Design", styles["FSDH3"]),
        _pdf_table(["Business Data", "Source / Contract", "UI Target", "Technical Rule"],
                   [[item.get("business_field") or item.get("mapping_id") or "Business data",
                     item.get("source") or "Approved OData / CDS",
                     item.get("ui_target") or "Fiori worklist / object page",
                     item.get("rule") or "GET for retrieval; POST/PATCH for permitted change"]
                    for item in (data.get("data_mappings") or [])]
                   or [["Approved FSD data", "OData service TBD", "Fiori application", "Backend remains authoritative"]],
                   [1.45, 1.65, 1.5, 1.9], styles, title="OData and data mapping"),
    ]
    story += _pdf_screenshot_placeholder(
        styles,
        "Figure: OData service metadata screenshot",
        "Insert a screenshot of the OData service metadata (SEGW, /IWFND/MAINT_SERVICE, or $metadata) for the approved entity set and operations.",
    )
    variants = data.get("process_steps") or []
    story += [
        _pdf_text("4.5.5 Workflow Technical Design", styles["FSDH3"]),
        _pdf_table(["Rule / Condition", "Technical Behavior", "Open Decision"],
                   [[
                       item.get("activity") or item.get("step_id"),
                       item.get("system_behavior") or item.get("outcome"),
                       "Confirm agent determination, thresholds and exception handling during design review"
                       if "approv" in _safe(item.get("activity")).lower() or "approv" in _safe(item.get("system_behavior")).lower()
                       else "None once configuration is approved",
                   ] for item in variants[:10]]
                   or [["Workflow (if required by FSD)", "Use standard SAP workflow / My Inbox when the FSD specifies approval", "Confirm during design review"]],
                   [1.75, 2.65, 2.1], styles, title="Workflow technical design"),
    ]
    story += _pdf_screenshot_placeholder(
        styles,
        "Figure: Flexible Workflow / My Inbox screenshot",
        "Insert a screenshot of the Flexible Workflow scenario or My Inbox task configuration once agent determination and thresholds are confirmed.",
    )
    story += [
        _pdf_text("4.5.6 Status Model", styles["FSDH3"]),
        _pdf_table(["Status", "Source", "UI Semantic State", "Definition"],
                   [[item.get("status") or item.get("name") or item.get("status_id"),
                     item.get("source") or item.get("domain") or "SAP backend",
                     item.get("semantic_state") or item.get("ui_state") or "TBD",
                     item.get("definition") or item.get("description") or "Confirm during design review"]
                    for item in (data.get("status_definitions") or [])]
                   or [["Draft / In Process / Complete", "SAP backend", "None / Information / Success", "Business-facing states derived from the FSD"]],
                   [1.4, 1.3, 1.4, 2.4], styles, title="Status model"),
        _pdf_text("4.5.7 Performance and Monitoring Design", styles["FSDH3"]),
        *_pdf_bullets([_safe(item.get("requirement") or item.get("description") or item) for item in data.get("nfr") or []]
                      or ["Use server-side filter, sort, page and select."], styles),
        _pdf_text("5. Testing Requirements", styles["FSDH1"]),
        _pdf_text("5.1 Key Unit / Assembly Test Conditions", styles["FSDH2"]),
        _pdf_table(["ID", "Condition / Test", "Expected Result", "Req Ref."],
                   [[item.get("id"), item.get("condition"), item.get("expected_result"), item.get("requirement_key") or item.get("cycle_ref")] for item in data.get("test_conditions", [])],
                   [.65, 2.25, 2.15, 1.45], styles, title="Unit and assembly test conditions"),
    ]
    story += _pdf_screenshot_placeholder(
        styles,
        "Figure: Unit / assembly test evidence screenshot",
        "Insert screenshots of unit or assembly test evidence (ABAP Unit, Gateway client, or Fiori test run) for the critical path covered above.",
    )
    story += [
        _pdf_text("6. Outstanding Issues", styles["FSDH1"]),
        _pdf_table(["Issue No.", "Description", "Assigned To", "Status", "Impact", "Resolution / Decision Needed"],
                   [[item.get("issue_no"), item.get("description"), item.get("assigned_to"), item.get("status"), item.get("impact"), item.get("resolution")] for item in data.get("outstanding_issues", [])],
                   [.7, 1.7, 1.0, .7, .85, 1.55], styles, title="Outstanding issues"),
        _pdf_text("7. Appendix", styles["FSDH1"]),
        _pdf_text("7.1 Glossary of Terms", styles["FSDH2"]),
        _pdf_table(["Term", "Definition"], [[item.get("term"), item.get("definition")] for item in data.get("glossary", [])],
                   [1.65, 4.85], styles, title="Glossary of terms"),
        _pdf_text("7.2 Additional Supporting / Reference Documentation", styles["FSDH2"]),
        *_pdf_bullets([
            f"Functional Solution Design: {_safe(source.get('title'))}, Version {_safe(source.get('version'))}",
            "Approved BRD and generated requirement baseline",
            "Quality Test Pack, traceability matrix, and generated Fiori/ABAP starter-code package",
        ], styles),
        _pdf_text("7.3 Key Requirement Traceability", styles["FSDH2"]),
        _pdf_table(["Design Area", "FSD Requirements Covered", "Technical Design Sections"],
                   [[item.get("capability"), item.get("requirement_ids"), "2.1, 3.2, 3.3, 4.5"] for item in data.get("traceability") or data.get("capability_map", [])],
                   [1.85, 2.45, 2.2], styles, title="Requirement traceability"),
        _pdf_text("7.4 Technical Design Review Notes", styles["FSDH2"]),
        *_pdf_bullets(data.get("review_notes", []), styles),
    ]
    pdf = _FSDTemplate(
        str(output), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=1.0 * inch, bottomMargin=.7 * inch,
        title=_safe(data.get("title")), author="SAP Project Copilot", project_name=project_name,
        document_label="SAP Technical Design Document", logo_left=images.get("logo_left"), logo_right=images.get("logo_right"),
    )
    pdf.multiBuild(story, onFirstPage=lambda canvas, doc: None, onLaterPages=lambda canvas, doc: None)
    _overlay_pdf_chrome(output, project_name, images.get("logo_left"), images.get("logo_right"), document_label="SAP Technical Design Document")


def export_technical_design(document_data: dict, project_name: str, project_id: str, artifact_root: str) -> dict[str, str]:
    display_name = _apply_tdd_identity(document_data, project_name)
    output_dir = (Path(artifact_root) / project_id).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{_slug(display_name)}-technical-design"
    docx_path = output_dir / f"{stem}.docx"
    pdf_path = output_dir / f"{stem}.pdf"
    left_logo = output_dir / f"{stem}-logo-left.png"
    right_logo = output_dir / f"{stem}-logo-right.png"
    process_image = output_dir / f"{stem}-flow.png"
    legend_image = output_dir / f"{stem}-flow-legend.png"

    _draw_logos(left_logo, right_logo)
    _draw_process_flow(document_data.get("process_steps") or [], process_image, "Technical Flow Diagram")
    _draw_process_legend(legend_image, steps=document_data.get("process_steps") or [])
    images = {"process": process_image, "legend": legend_image, "logo_left": left_logo, "logo_right": right_logo}

    doc = WordDocument()
    _style_word(doc, display_name, {"left": left_logo, "right": right_logo}, document_label="SAP Technical Design Document")
    _cover(doc, document_data, display_name)
    _add_contents(doc, document_data, images)
    doc.save(str(docx_path))
    _build_tdd_pdf(document_data, images, pdf_path, display_name)
    return {"docx": str(docx_path), "pdf": str(pdf_path)}
