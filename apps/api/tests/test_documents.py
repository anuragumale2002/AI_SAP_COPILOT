import io

from docx import Document

from app.services.documents import extract_document


def test_docx_extraction_includes_tables_and_headings():
    document = Document()
    document.add_heading("Business Requirements", level=1)
    document.add_paragraph("The solution shall provide a consolidated supplier view for authorized users.")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Requirement ID"
    table.cell(0, 1).text = "Requirement"
    table.cell(0, 2).text = "Priority"
    table.cell(1, 0).text = "BR-001"
    table.cell(1, 1).text = "Users shall search suppliers by Supplier ID and name."
    table.cell(1, 2).text = "Must"

    buffer = io.BytesIO()
    document.save(buffer)

    page_count, chunks = extract_document("supplier.docx", buffer.getvalue())
    combined = "\n".join(chunk.text for chunk in chunks)
    locators = [chunk.locator for chunk in chunks]

    assert page_count == 1
    assert "consolidated supplier view" in combined
    assert "BR-001" in combined
    assert "Users shall search suppliers" in combined
    assert any(locator.startswith("Table 1, row") for locator in locators)
