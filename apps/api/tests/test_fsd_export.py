from pathlib import Path

from docx import Document
from PIL import Image

from app.services.fsd_export import _draw_process_flow, _pdf_styles, _pdf_table, _pdf_table_cell, _pdf_table_number, _word_table


def test_long_generated_table_row_fits_on_a_pdf_page():
    widths = [1.0, 1.25, 2.0, 1.5, 0.75]
    very_long_detail = "A generated process-flow explanation " * 800
    table = _pdf_table(
        ["Flow Type", "Trigger", "Flow", "Outcome", "Req IDs"],
        [["Exception", very_long_detail, very_long_detail, very_long_detail, list(range(1, 118))]],
        widths,
        _pdf_styles(),
    )

    _, height = table.wrap(468, 646)

    assert height < 646
    assert "Full detail is in the DOCX" in _pdf_table_cell(very_long_detail, 1.25)


def test_tables_receive_sequential_captions_in_word_and_pdf():
    document = Document()
    _word_table(document, ["Field", "Detail"], [["A", "B"]], [1.0, 1.0])
    _word_table(document, ["Role", "Name"], [["Owner", "TBD"]], [1.0, 1.0])

    assert document.paragraphs[0].text == "Table 1: Record details"
    assert document.paragraphs[1].text == "Table 2: Role summary"

    _pdf_table_number.set(0)
    table = _pdf_table(["Field", "Detail"], [["A", "B"]], [1.0, 1.0], _pdf_styles())
    assert table._cellvalues[-1][0].getPlainText() == "Table 1: Record details"


def test_process_flow_uses_role_swimlanes(tmp_path: Path):
    output = tmp_path / "flow.png"
    steps = [
        {"step_id": "S1", "actor": "Requester", "activity": "Submit request", "decision_or_rule": "TBD"},
        {"step_id": "S2", "actor": "Approver", "activity": "Review request", "decision_or_rule": "Approved?"},
    ]

    _draw_process_flow(steps, output)

    with Image.open(output) as image:
        assert image.width == 1900
        assert image.height >= 545
